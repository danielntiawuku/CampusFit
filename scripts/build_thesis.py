#!/usr/bin/env python3
"""
Build the CampusFit thesis from the chapter markdown files into Word documents.

Default (full thesis) writes dee/CampusFit_Thesis.docx with:
  * title page, abstract, an updatable table-of-contents field,
  * Times New Roman 12pt, 1.5 spacing, justified body text,
  * Heading 1/2/3 mapped to chapters, sections and subsections,
  * bullet/numbered lists, markdown tables, bold and italic runs.

--split additionally writes one document per chapter into dee/thesis_split/:
  CampusFit_Chapter_1.docx ... CampusFit_Chapter_6.docx
Each split document carries the thesis title as a preface and the chapter
exactly as it appears in the full document (same styles, same content).

Usage:
  python scripts/build_thesis.py                # full thesis only
  python scripts/build_thesis.py --split        # full thesis + 6 chapter files
  python scripts/build_thesis.py --split-only   # 6 chapter files only
  python scripts/build_thesis.py --out path.docx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

ROOT = Path(__file__).resolve().parent.parent
THESIS_TITLE = (
    "CAMPUSFIT: DESIGN AND IMPLEMENTATION OF A GAMIFIED, CHECKPOINT-BASED "
    "MOBILE FITNESS APPLICATION FOR UNIVERSITY STUDENTS"
)
CHAPTERS = [
    "00_front.md",
    "chapter1.md",
    "chapter2.md",
    "chapter2b.md",
    "chapter3.md",
    "chapter4.md",
    "chapter5.md",
    "chapter6.md",
    "references.md",
    "appendices.md",
]
# (files, output name, chapter label shown on the split title block)
SPLITS = [
    (["chapter1.md"], "CampusFit_Chapter_1.docx", "CHAPTER ONE: INTRODUCTION"),
    (["chapter2.md", "chapter2b.md"], "CampusFit_Chapter_2.docx",
     "CHAPTER TWO: THEORY, BACKGROUND AND REVIEW"),
    (["chapter3.md"], "CampusFit_Chapter_3.docx", "CHAPTER THREE: METHODOLOGY"),
    (["chapter4.md"], "CampusFit_Chapter_4.docx",
     "CHAPTER FOUR: DESIGN AND IMPLEMENTATION (CONTRIBUTION)"),
    (["chapter5.md"], "CampusFit_Chapter_5.docx", "CHAPTER FIVE: RESEARCH EVALUATION"),
    (["chapter6.md"], "CampusFit_Chapter_6.docx",
     "CHAPTER SIX: CONCLUSION AND FUTURE WORKS"),
]

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_inline(paragraph, text: str) -> None:
    """Render **bold** and *italic* runs inside a paragraph."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            paragraph.add_run(chunk)


def add_toc_field(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Table of contents — open in Word and press F9 (or right-click → Update Field) to generate."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading_specs = {
        "Heading 1": (16, True, 18, 12),
        "Heading 2": (14, True, 14, 8),
        "Heading 3": (12, True, 12, 6),
    }
    for name, (size, bold, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0x18, 0x1C, 0x1A)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.line_spacing = 1.5

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) or not c.strip() for c in r)]
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            text = row[j].strip() if j < len(row) else ""
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(text)
            run.font.size = Pt(10.5)
            run.font.name = "Times New Roman"
            if i == 0:
                run.bold = True
    doc.add_paragraph()


def add_preface(doc: Document, lines: list[tuple[str, int, bool]]) -> None:
    """Centred title block used by the split chapter documents."""
    first = True
    for text, size, bold in lines:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(72 if first else 18)
        paragraph.paragraph_format.space_after = Pt(12)
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        first = False


def build_doc(out_path: Path, chapters: list[str], preface: list[tuple[str, int, bool]] | None = None) -> None:
    doc = Document()
    configure_styles(doc)

    if preface:
        add_preface(doc, preface)

    # Title-page treatment only for the standalone full thesis (no preface).
    consume_title = preface is None
    h1_count = 0

    for filename in chapters:
        path = ROOT / "thesis" / filename
        if not path.exists():
            raise SystemExit(f"missing chapter: {path}")
        lines = path.read_text(encoding="utf-8").splitlines()

        index = 0
        while index < len(lines):
            line = lines[index].rstrip()
            index += 1

            if not line.strip() or line.strip() == "---":
                continue

            # markdown table
            if line.lstrip().startswith("|"):
                block = []
                while index < len(lines) and lines[index].lstrip().startswith("|"):
                    block.append([c for c in lines[index].strip().strip("|").split("|")])
                    index += 1
                add_table(doc, block)
                continue

            if line.startswith("# "):
                text = line[2:].strip()
                if consume_title:
                    # Title page: centered, larger, no TOC entry
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(72)
                    run = paragraph.add_run(text)
                    run.bold = True
                    run.font.size = Pt(20)
                    run.font.name = "Times New Roman"
                    consume_title = False
                    continue
                paragraph = doc.add_paragraph(style="Heading 1")
                if h1_count > 0 or preface is None:
                    paragraph.paragraph_format.page_break_before = True
                add_inline(paragraph, text)
                h1_count += 1
                continue

            if line.startswith("## "):
                if filename == "00_front.md":
                    paragraph = doc.add_paragraph(style="Heading 1")
                    add_inline(paragraph, line[3:].strip())
                else:
                    paragraph = doc.add_paragraph(style="Heading 2")
                    add_inline(paragraph, line[3:].strip())
                if line[3:].strip().upper().startswith("TABLE OF CONTENTS"):
                    add_toc_field(doc)
                continue

            if line.startswith("### "):
                paragraph = doc.add_paragraph(style="Heading 3")
                add_inline(paragraph, line[4:].strip())
                continue

            if re.match(r"^- ", line):
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.line_spacing = 1.5
                add_inline(paragraph, line[2:].strip())
                continue

            numbered = re.match(r"^(\d+)\.\s+(.*)", line)
            if numbered:
                paragraph = doc.add_paragraph(style="List Number")
                paragraph.paragraph_format.line_spacing = 1.5
                add_inline(paragraph, numbered.group(2))
                continue

            paragraph = doc.add_paragraph()
            add_inline(paragraph, line.strip())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"wrote {out_path}")


def build(out_path: Path) -> None:
    build_doc(out_path, CHAPTERS)


def build_split(split_dir: Path) -> None:
    for files, name, label in SPLITS:
        build_doc(
            split_dir / name,
            files,
            preface=[
                (THESIS_TITLE, 12, True),
                (label, 16, True),
            ],
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(ROOT / "CampusFit_Thesis.docx"))
    parser.add_argument("--split", action="store_true", help="also write one docx per chapter")
    parser.add_argument("--split-only", action="store_true", help="only write the per-chapter docx files")
    parser.add_argument("--split-dir", default=str(ROOT / "thesis_split"))
    args = parser.parse_args()

    if not args.split_only:
        build(Path(args.out))
    if args.split or args.split_only:
        build_split(Path(args.split_dir))


if __name__ == "__main__":
    main()
